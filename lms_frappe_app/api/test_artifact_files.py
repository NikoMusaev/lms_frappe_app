# Copyright (c) 2026, NikoMusaev and Contributors
# See license.txt

import io

import frappe
from frappe.tests import IntegrationTestCase

from lms_frappe_app.agent_learning import artifact_files, course_builder
from lms_frappe_app.api import authoring, manager, student
from lms_frappe_app.tests.sample_data import (
	добавить_в_организацию,
	зачислить,
	создать_куратора,
	создать_менеджера,
	создать_организацию,
	создать_ученика,
	создать_урок,
)

CSV = "Месяц;Выручка;Расходы\nЯнварь;100;80\nФевраль;120;90\n".encode()


def xlsx(с_формулой: bool = True) -> bytes:
	from openpyxl import Workbook

	книга = Workbook()
	лист = книга.active
	лист.append(["Месяц", "Выручка", "Расходы", "Прибыль"])
	лист.append(["Январь", 100, 80, "=B2-C2" if с_формулой else 20])
	поток = io.BytesIO()
	книга.save(поток)
	return поток.getvalue()


class IntegrationTestArtifactFiles(IntegrationTestCase):
	"""Блок документа как файл или ссылка (learning-services#315)."""

	def setUp(self):
		self.addCleanup(frappe.set_user, "Administrator")
		суффикс = frappe.generate_hash(length=6)
		self.ученик = создать_ученика(f"file-{суффикс}@example.com")
		self.урок = создать_урок(f"Урок {суффикс}")
		self.курс = зачислить(self.ученик, self.урок)
		frappe.get_doc(
			{
				"doctype": "Agent Course Artifact",
				"course": self.курс,
				"slug": "plan",
				"title": "План",
				"blocks": [
					{"block_key": "money", "title": "Финплан", "kind": "file", "accept": "xlsx,csv", "lesson": self.урок},
					{"block_key": "crm", "title": "База клиентов", "kind": "link"},
					{"block_key": "goal", "title": "Цель"},
				],
			}
		).insert(ignore_permissions=True)
		frappe.set_user(self.ученик)

	def загрузить(self, имя: str = "plan.csv", данные: bytes = CSV, ключ: str = "money") -> dict:
		return student.контракт(student._положить_файл)(self.ученик, self.курс, "plan", ключ, имя, данные)

	def запись_файла(self, ключ: str = "money") -> str | None:
		"""Имя `File` в строке блока. По адресу искать нельзя: Frappe отдаёт
		одинаковому содержимому один `file_url` на несколько записей."""
		документ = frappe.db.get_value("Agent Student Artifact", {"student": self.ученик, "course": self.курс})
		return frappe.db.get_value(
			"Agent Artifact Content", {"parent": документ, "block_key": ключ}, "file"
		)

	def блок(self, ключ: str) -> dict:
		блоки = student.artifact(self.курс, "plan")["data"]["blocks"]
		return next(б for б in блоки if б["key"] == ключ)

	# --- файл ---

	def test_загруженный_csv_виден_файлом_и_срезом(self):
		ответ = self.загрузить()["data"]

		self.assertEqual(ответ["file"]["name"], "plan.csv")
		блок = self.блок("money")
		self.assertEqual(блок["kind"], "file")
		self.assertEqual(блок["accept"], ["xlsx", "csv"])
		self.assertEqual(блок["file"]["type"], "csv")
		self.assertIn("| Месяц | Выручка | Расходы |", блок["preview"])
		self.assertIn("| Февраль | 120 | 90 |", блок["preview"])
		перечень = student.artifact(self.курс)["data"]["artifacts"][0]
		self.assertEqual(перечень["blocks_filled"], 1, "файл без текста — блок заполнен")

	def test_срез_xlsx_показывает_формулу_без_значения(self):
		self.загрузить("plan.xlsx", xlsx())

		срез = self.блок("money")["preview"]

		self.assertIn("| Месяц | Выручка | Расходы | Прибыль |", срез)
		self.assertIn("=B2-C2", срез, "у собранного программой файла значения нет — видна формула")

	def test_пустые_колонки_справа_в_срез_не_попадают(self):
		from openpyxl import Workbook
		from openpyxl.styles import Font

		книга = Workbook()
		лист = книга.active
		лист.append(["Гипотеза", "Охват"])
		лист.append(["Кофейня", 100])
		# Отформатированная, но пустая колонка далеко справа — как в живом файле.
		лист.cell(row=1, column=25).font = Font(bold=True)
		поток = io.BytesIO()
		книга.save(поток)

		self.загрузить("rice.xlsx", поток.getvalue())
		срез = self.блок("money")["preview"]

		self.assertEqual(срез.splitlines()[0], "| Гипотеза | Охват |")

	def test_новый_файл_замещает_прежний(self):
		self.загрузить()
		первый = self.запись_файла()
		второй = self.загрузить("plan.xlsx", xlsx(с_формулой=False))["data"]["file"]

		self.assertEqual(self.блок("money")["file"]["name"], "plan.xlsx")
		self.assertFalse(frappe.db.exists("File", первый), "прежний файл удалён")
		self.assertTrue(второй["url"].startswith("/private/files/"), "файл приватный")

	def test_файл_не_того_типа_отказ(self):
		ответ = self.загрузить("photo.png", b"\x89PNG")

		self.assertEqual(ответ["error"]["code"], artifact_files.ФАЙЛ_НЕ_ТОГО_ТИПА)
		self.assertEqual(ответ["error"]["accept"], ["xlsx", "csv"])

	def test_пустой_файл_и_текстовый_блок_отказ(self):
		self.assertEqual(self.загрузить(данные=b"")["error"]["code"], artifact_files.ФАЙЛА_НЕТ)
		self.assertEqual(self.загрузить(ключ="goal")["error"]["code"], artifact_files.ВИД_НЕ_ТОТ)

	def test_слишком_большой_файл_отказ(self):
		frappe.set_user("Administrator")
		frappe.db.set_single_value("Agent Learning Settings", "artifact_file_max_mb", 1)
		frappe.clear_document_cache("Agent Learning Settings", "Agent Learning Settings")
		self.addCleanup(frappe.clear_document_cache, "Agent Learning Settings", "Agent Learning Settings")
		self.addCleanup(frappe.db.set_single_value, "Agent Learning Settings", "artifact_file_max_mb", 10)
		frappe.set_user(self.ученик)

		ответ = self.загрузить(данные=b"a;b\n" * 400_000)

		self.assertEqual(ответ["error"]["code"], artifact_files.ФАЙЛ_СЛИШКОМ_БОЛЬШОЙ)

	def test_очистка_удаляет_файл(self):
		self.загрузить()
		файл = self.запись_файла()

		student.update_artifact(self.курс, "plan", "money", clear=True)

		self.assertIsNone(self.блок("money")["file"])
		self.assertFalse(frappe.db.exists("File", файл))

	def test_start_lesson_приносит_файл_блока_урока(self):
		self.загрузить()

		блоки = student.start_lesson(lesson=self.урок)["data"]["artifact_blocks"]

		self.assertEqual([б["key"] for б in блоки], ["money"])
		self.assertEqual(блоки[0]["file"]["name"], "plan.csv")
		self.assertTrue(блоки[0]["preview"])

	def test_файл_видит_только_ученик(self):
		self.загрузить()
		файл = frappe.get_doc("File", self.запись_файла())
		frappe.set_user("Administrator")
		суффикс = frappe.generate_hash(length=6)
		другой = создать_ученика(f"other-{суффикс}@example.com")
		организация = создать_организацию(f"Компания {суффикс}")
		добавить_в_организацию(self.ученик, организация)
		руководитель = создать_менеджера(f"mg-{суффикс}@example.com", организация)

		self.assertTrue(frappe.has_permission("File", doc=файл, user=self.ученик))
		for чужой in (другой, руководитель):
			with self.subTest(чужой=чужой):
				self.assertFalse(frappe.has_permission("File", doc=файл, user=чужой))

	# --- ссылка ---

	def test_ссылка_записывается_в_блок_ссылку(self):
		ответ = student.update_artifact(self.курс, "plan", "crm", url="https://crm.example.com/base")

		self.assertTrue(ответ["ok"])
		self.assertEqual(self.блок("crm")["url"], "https://crm.example.com/base")
		self.assertEqual(ответ["data"]["blocks_filled"], 1)

	def test_ссылка_не_в_тот_блок_и_кривая_отказ(self):
		мимо = student.update_artifact(self.курс, "plan", "goal", url="https://example.com")
		кривая = student.update_artifact(self.курс, "plan", "crm", url="javascript:alert(1)")

		self.assertEqual(мимо["error"]["code"], artifact_files.ВИД_НЕ_ТОТ)
		self.assertEqual(кривая["error"]["code"], artifact_files.НЕВЕРНАЯ_ССЫЛКА)

	def test_текстовый_блок_как_прежде(self):
		student.update_artifact(self.курс, "plan", "goal", "Открыть вторую кофейню")

		блок = self.блок("goal")
		self.assertEqual((блок["kind"], блок["content"], блок["file"], блок["url"]), ("text", "Открыть вторую кофейню", None, None))

	# --- отчёт руководителя ---

	def test_отчёт_руководителя_считает_файл_заполненным(self):
		self.загрузить()
		frappe.set_user("Administrator")
		организация = создать_организацию(f"Отчёт {frappe.generate_hash(length=6)}")
		добавить_в_организацию(self.ученик, организация)
		frappe.get_doc(
			{"doctype": "Course Allocation", "organization": организация, "course": self.курс}
		).insert(ignore_permissions=True)
		руководитель = создать_менеджера(f"rep-{frappe.generate_hash(length=6)}@example.com", организация)
		frappe.set_user(руководитель)

		строка = next(с for с in manager.org_report()["data"]["rows"] if с["user"] == self.ученик)

		self.assertEqual(строка["document"], {"blocks_total": 3, "blocks_filled": 1})


class IntegrationTestArtifactKindsAuthoring(IntegrationTestCase):
	"""Вид блока в авторинге и готовности курса."""

	def setUp(self):
		self.addCleanup(frappe.set_user, "Administrator")
		суффикс = frappe.generate_hash(length=6)
		self.урок = создать_урок(f"Урок {суффикс}")
		self.курс = frappe.db.get_value(
			"Course Chapter", frappe.db.get_value("Course Lesson", self.урок, "chapter"), "course"
		)
		frappe.set_user(создать_куратора(f"au-{суффикс}@example.com"))

	def схема(self, блоки: list[dict]) -> dict:
		return authoring.set_course_artifact(self.курс, "plan", "План", блоки)

	def test_вид_и_форматы_сохраняются_и_видны_в_черновике(self):
		self.схема([{"key": "money", "title": "Финплан", "kind": "file", "accept": ".XLSX, csv"}])

		блок = authoring.course_draft(self.курс)["data"]["artifacts"][0]["blocks"][0]

		self.assertEqual((блок["kind"], блок["accept"]), ("file", ["xlsx", "csv"]))

	def test_неизвестный_вид_отказ(self):
		ответ = self.схема([{"key": "money", "title": "Финплан", "kind": "pdf"}])

		self.assertEqual(ответ["error"]["code"], authoring.НЕВЕРНЫЙ_ВИД_БЛОКА)

	def test_готовность_предупреждает_о_файле_без_форматов(self):
		self.схема([{"key": "money", "title": "Финплан", "kind": "file"}])

		коды = [п["code"] for п in course_builder.проверить_готовность(self.курс)["warnings"]]

		self.assertIn(course_builder.ФАЙЛ_БЕЗ_ФОРМАТОВ, коды)
