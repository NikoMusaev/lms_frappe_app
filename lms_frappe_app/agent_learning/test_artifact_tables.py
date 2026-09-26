# Copyright (c) 2026, NikoMusaev and contributors
# For license information, please see license.txt

"""Поля и таблицы документа без базы (learning-services#330).

Схема — реестр рисков из курса «Риски проекта» в миниатюре: блок рисков
заводит строки, блок оценки дописывает им баллы, ранг и отметку «в работе»
считает сервер, ответ обязателен только у рисков в работе.
"""

import unittest

from lms_frappe_app.agent_learning import artifact_tables as т
from lms_frappe_app.agent_learning.errors import Отказ

ЦЕЛИ = {
	"block_key": "goals",
	"spec": {
		"table": "goals",
		"prefix": "G",
		"columns": [
			{"key": "goal", "title": "Цель", "type": "text", "required": True},
			{"key": "measure", "title": "Мера отклонения", "type": "text"},
		],
	},
}
РИСКИ = {
	"block_key": "risks",
	"spec": {
		"table": "register",
		"prefix": "R",
		"title": "Реестр",
		"columns": [
			{
				"key": "kind",
				"title": "Вид",
				"type": "select",
				"options": ["угроза", "возможность"],
				"required": True,
			},
			{"key": "event", "title": "Событие", "type": "text", "required": True},
			{"key": "goal_ref", "title": "Цель", "type": "ref", "ref": "goals"},
		],
		"views": [{"type": "matrix", "x": "impact", "y": "probability", "highlight": "in_work"}],
	},
}
ШКАЛА = {
	"block_key": "scales",
	"spec": {
		"table": "probability_scale",
		"prefix": "P",
		"columns": [
			{"key": "score", "title": "Балл", "type": "number"},
			{"key": "level", "title": "Уровень", "type": "text"},
		],
		"rows": [{"score": 1, "level": "Редко"}, {"score": 5, "level": "Почти наверняка"}],
	},
}
ОЦЕНКА = {
	"block_key": "assessment",
	"spec": {
		"table": "register",
		"fields": [{"key": "threshold", "title": "Порог внимания", "type": "number", "required": True}],
		"columns": [
			{
				"key": "probability",
				"title": "Вероятность",
				"type": "scale",
				"labels": "probability_scale",
				"required": True,
			},
			{"key": "impact", "title": "Влияние", "type": "scale", "required": True},
			{"key": "rank", "title": "Ранг", "type": "formula", "formula": "probability * impact"},
			{"key": "in_work", "title": "В работе", "type": "formula", "formula": "rank >= threshold"},
		],
	},
}
ОТВЕТЫ = {
	"block_key": "responses",
	"spec": {
		"table": "register",
		"columns": [
			{"key": "measure_text", "title": "Мера", "type": "text", "required": "in_work"},
			{"key": "due", "title": "Срок", "type": "date"},
			{"key": "sponsor", "title": "Спонсору", "type": "check", "max": 1},
		],
	},
}
ТЕКСТ = {"block_key": "worries", "spec": None}


def схема():
	блоки = [dict(б) for б in (ТЕКСТ, ЦЕЛИ, РИСКИ, ШКАЛА, ОЦЕНКА, ОТВЕТЫ)]
	for б in блоки:
		б["spec"] = т.проверить_спек(б["spec"], б["block_key"])
	т.проверить_документ(блоки)
	return блоки


class TestСхема(unittest.TestCase):
	def test_таблица_собирает_колонки_всех_блоков_и_хозяина(self):
		таблица = т.таблицы_схемы(схема())["register"]
		self.assertEqual(таблица["owner"], "risks")
		self.assertEqual(
			[к["key"] for к in таблица["columns"]],
			[
				"kind",
				"event",
				"goal_ref",
				"probability",
				"impact",
				"rank",
				"in_work",
				"measure_text",
				"due",
				"sponsor",
			],
		)
		self.assertEqual(таблица["columns"][3]["block"], "assessment")

	def test_формула_на_неизвестное_имя_отказ(self):
		блоки = схема()
		блоки[4]["spec"]["columns"][2]["formula"] = "probability * weight"
		with self.assertRaises(Отказ) as отказ:
			т.проверить_документ(блоки)
		self.assertEqual(отказ.exception.код, т.НЕВЕРНАЯ_СХЕМА)

	def test_ссылка_на_таблицу_которой_нет_отказ(self):
		блоки = схема()
		блоки[2]["spec"]["columns"][2]["ref"] = "nope"
		with self.assertRaises(Отказ):
			т.проверить_документ(блоки)

	def test_выбор_без_вариантов_отказ(self):
		with self.assertRaises(Отказ):
			т.проверить_спек({"columns": [{"key": "kind", "type": "select"}]}, "x")

	def test_колонка_id_занята(self):
		with self.assertRaises(Отказ):
			т.проверить_спек({"columns": [{"key": "id"}]}, "x")


class TestФормулы(unittest.TestCase):
	def test_арифметика_и_сравнения(self):
		self.assertEqual(т.вычислить(т.разобрать("a * b + 1"), {"a": 3, "b": 4}), 13)
		self.assertTrue(т.вычислить(т.разобрать("(a - 1) >= 2 and not b"), {"a": 3, "b": False}))

	def test_пустой_операнд_даёт_пусто(self):
		self.assertIsNone(т.вычислить(т.разобрать("a * b"), {"a": 3}))

	def test_деление_на_ноль_пусто(self):
		self.assertIsNone(т.вычислить(т.разобрать("a / b"), {"a": 3, "b": 0}))

	def test_мусор_не_разбирается(self):
		for формула in ("a *", "__import__('os')", "a ** b", "(a"):
			with self.assertRaises(Отказ, msg=формула):
				т.разобрать(формула)


class TestЗапись(unittest.TestCase):
	def setUp(self):
		self.блоки = схема()
		д = т.записать(self.блоки, "goals", {}, rows=[{"goal": "Запуск до 1 марта"}])
		self.д = т.записать(
			self.блоки,
			"risks",
			д,
			rows=[
				{"kind": "Угроза", "event": "Подрядчик уйдёт на другой объект", "goal_ref": "g1"},
				{"kind": "возможность", "event": "Готовый модуль сэкономит месяц"},
			],
		)

	def test_хозяин_заводит_строки_с_номерами(self):
		строки = self.д["tables"]["register"]
		self.assertEqual([с["id"] for с in строки], ["R1", "R2"])
		# Выбор и ссылка — к каноническому написанию.
		self.assertEqual(строки[0]["kind"], "угроза")
		self.assertEqual(строки[0]["goal_ref"], "G1")

	def test_чужой_блок_дописывает_колонки_но_не_заводит_строки(self):
		д = т.записать(self.блоки, "assessment", self.д, rows=[{"id": "R1", "probability": "4", "impact": 5}])
		self.assertEqual(д["tables"]["register"][0]["probability"], 4)
		with self.assertRaises(Отказ) as отказ:
			т.записать(self.блоки, "assessment", self.д, rows=[{"probability": 4}])
		self.assertEqual(отказ.exception.код, т.СТРОКИ_НЕ_ЗАВОДИТ)

	def test_чужую_колонку_не_записать(self):
		with self.assertRaises(Отказ) as отказ:
			т.записать(self.блоки, "assessment", self.д, rows=[{"id": "R1", "event": "другое"}])
		self.assertEqual(отказ.exception.код, т.НЕТ_КОЛОНКИ)

	def test_формулу_не_записать(self):
		with self.assertRaises(Отказ):
			т.записать(self.блоки, "assessment", self.д, rows=[{"id": "R1", "rank": 20}])

	def test_балл_вне_шкалы_и_неизвестный_вариант_отказ(self):
		with self.assertRaises(Отказ):
			т.записать(self.блоки, "assessment", self.д, rows=[{"id": "R1", "probability": 6}])
		with self.assertRaises(Отказ):
			т.записать(self.блоки, "risks", self.д, rows=[{"id": "R1", "kind": "проблема"}])

	def test_ссылка_на_несуществующую_строку_отказ(self):
		with self.assertRaises(Отказ):
			т.записать(self.блоки, "risks", self.д, rows=[{"id": "R1", "goal_ref": "G9"}])

	def test_дата_по_русски_приводится_к_iso(self):
		д = т.записать(self.блоки, "responses", self.д, rows=[{"id": "R1", "due": "05.03.2027"}])
		self.assertEqual(д["tables"]["register"][0]["due"], "2027-03-05")

	def test_флажков_не_больше_предела(self):
		д = т.записать(self.блоки, "responses", self.д, rows=[{"id": "R1", "sponsor": True}])
		with self.assertRaises(Отказ) as отказ:
			т.записать(self.блоки, "responses", д, rows=[{"id": "R2", "sponsor": True}])
		self.assertEqual(отказ.exception.код, т.ФЛАЖКОВ_СЛИШКОМ_МНОГО)

	def test_удалённый_номер_не_возвращается(self):
		д = т.записать(self.блоки, "risks", self.д, delete_rows=["R2"])
		д = т.записать(self.блоки, "risks", д, rows=[{"kind": "угроза", "event": "Сезон начнётся раньше"}])
		self.assertEqual([с["id"] for с in д["tables"]["register"]], ["R1", "R3"])

	def test_удаляет_только_хозяин(self):
		with self.assertRaises(Отказ):
			т.записать(self.блоки, "assessment", self.д, delete_rows=["R1"])

	def test_поле_блока_и_чужое_поле(self):
		д = т.записать(self.блоки, "assessment", self.д, fields={"threshold": "12"})
		self.assertEqual(д["fields"]["threshold"], 12)
		with self.assertRaises(Отказ) as отказ:
			т.записать(self.блоки, "risks", self.д, fields={"threshold": 12})
		self.assertEqual(отказ.exception.код, т.НЕТ_ПОЛЯ)


class TestВычисленияИЗаполненность(unittest.TestCase):
	def setUp(self):
		self.блоки = схема()
		д = т.записать(
			self.блоки,
			"risks",
			{},
			rows=[{"kind": "угроза", "event": "A"}, {"kind": "угроза", "event": "B"}],
		)
		д = т.записать(self.блоки, "assessment", д, fields={"threshold": 12})
		self.д = т.записать(
			self.блоки,
			"assessment",
			д,
			rows=[{"id": "R1", "probability": 4, "impact": 4}, {"id": "R2", "probability": 2, "impact": 3}],
		)

	def блок(self, ключ):
		return next(б for б in self.блоки if б["block_key"] == ключ)

	def test_ранг_и_в_работе_считаются(self):
		ряды = т.вычисленные(т.таблицы_схемы(self.блоки)["register"], self.д)
		self.assertEqual([(р["rank"], р["in_work"]) for р in ряды], [(16, True), (6, False)])

	def test_ответ_обязателен_только_у_рисков_в_работе(self):
		ответы = self.блок("responses")
		self.assertFalse(т.заполнен(ответы, self.д, self.блоки))
		self.assertEqual(
			т.пустые_клетки(ответы, self.д, self.блоки), [{"row": "R1", "column": "measure_text"}]
		)
		д = т.записать(
			self.блоки, "responses", self.д, rows=[{"id": "R1", "measure_text": "уточнить у Анны"}]
		)
		self.assertTrue(т.заполнен(ответы, д, self.блоки))

	def test_оценка_заполнена_когда_поле_и_колонки_есть(self):
		self.assertTrue(т.заполнен(self.блок("assessment"), self.д, self.блоки))
		д = т.записать(self.блоки, "risks", self.д, rows=[{"kind": "угроза", "event": "C"}])
		self.assertFalse(т.заполнен(self.блок("assessment"), д, self.блоки))

	def test_прежний_текст_считается_пока_значений_нет(self):
		цели = self.блок("goals")
		self.assertTrue(т.заполнен(цели, self.д, self.блоки, "| G1 | запуск |"))
		self.assertFalse(т.заполнен(цели, self.д, self.блоки, ""))

	def test_текстовый_блок_не_касается_таблиц(self):
		self.assertIsNone(т.заполнен(self.блок("worries"), self.д, self.блоки, "тревоги"))

	def test_нетронутая_заготовка_не_заполняет_блок(self):
		шкала = self.блок("scales")
		self.assertFalse(т.заполнен(шкала, self.д, self.блоки))
		д = т.записать(self.блоки, "scales", self.д, rows=[{"id": "P1", "level": "На памяти не было"}])
		self.assertTrue(т.заполнен(шкала, д, self.блоки))

	def test_заготовка_строк_видна_до_первой_правки(self):
		шкала = т.таблицы_документа(self.блоки, self.д)["probability_scale"]
		self.assertEqual([р["id"] for р in шкала["rows"]], ["P1", "P2"])
		д = т.записать(self.блоки, "scales", self.д, rows=[{"id": "P2", "level": "Наверняка"}])
		self.assertEqual(д["tables"]["probability_scale"][1], {"id": "P2", "score": 5, "level": "Наверняка"})

	def test_markdown_для_агента(self):
		текст = т.таблицы_документа(self.блоки, self.д)["register"]["markdown"]
		self.assertIn("| ID | Вид | Событие |", текст)
		self.assertIn("| R1 | угроза | A |  | 4 | 4 | 16 | да |", текст)

	def test_xlsx_держит_формулы(self):
		from io import BytesIO

		from openpyxl import load_workbook

		книга = load_workbook(BytesIO(т.книга_xlsx("Реестр", self.блоки, self.д)))
		self.assertEqual(книга.sheetnames[0], "Шапка")
		self.assertIn("Реестр", книга.sheetnames)
		лист = книга["Реестр"]
		шапка = [к.value for к in лист[1]]
		ранг = шапка.index("Ранг")
		self.assertEqual(лист.cell(row=2, column=ранг + 1).value, "=E2*F2")
		в_работе = лист.cell(row=2, column=шапка.index("В работе") + 1).value
		self.assertEqual(в_работе, "=G2>='Шапка'!$B$2")


if __name__ == "__main__":
	unittest.main()
