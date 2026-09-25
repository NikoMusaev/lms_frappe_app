# Copyright (c) 2026, NikoMusaev and contributors
# For license information, please see license.txt

"""Блок документа курса как файл или ссылка (learning-services#258, #315).

Документ курса хранит блоки как markdown. Но часть блоков в жизни не текст:
финансовый план — таблица с формулами, журнал обращений пополняется каждый
день, база клиентов живёт в CRM. Такой блок объявляется файлом или ссылкой, и
рядом с текстом (снимком, выводами) у него есть:

- **файл** — приватный `File`, привязанный к документу ученика. Права на него
  Frappe берёт у документа: видит только ученик;
- **ссылка** — адрес внешнего документа; не проверяется, это указатель;
- **срез** — шапка и первые строки таблицы markdown-текстом. Считается при
  загрузке: по нему агент сверяет готовность блока, не открывая файл.
"""

from __future__ import annotations

import csv
import io
from urllib.parse import urlparse

import frappe

from lms_frappe_app.agent_learning.doctype.agent_learning_settings.agent_learning_settings import (
	настройка,
)
from lms_frappe_app.agent_learning.errors import Отказ

ТЕКСТ, ФАЙЛ, ССЫЛКА = "text", "file", "link"
ВИДЫ = (ТЕКСТ, ФАЙЛ, ССЫЛКА)

ВИД_НЕ_ТОТ = "artifact_kind_mismatch"
ФАЙЛ_НЕ_ТОГО_ТИПА = "artifact_file_type"
ФАЙЛ_СЛИШКОМ_БОЛЬШОЙ = "artifact_file_too_large"
ФАЙЛА_НЕТ = "artifact_file_missing"
НЕВЕРНАЯ_ССЫЛКА = "artifact_invalid_url"

#: Из чего строится срез. Остальные типы хранятся без него.
ТАБЛИЦЫ = ("csv", "xlsx")
#: Ширина ячейки в срезе: длинный текст режется, иначе одна ячейка съедает
#: строку таблицы целиком.
ШИРИНА_ЯЧЕЙКИ = 60


def вид(блок) -> str:
	"""Вид блока схемы; старые блоки без поля — текст."""
	значение = блок.get("kind") if isinstance(блок, dict) else getattr(блок, "kind", None)
	return значение or ТЕКСТ


def допустимые(блок) -> list[str]:
	"""Расширения, которые принимает блок-файл, без точки и в нижнем регистре."""
	строка = блок.get("accept") if isinstance(блок, dict) else getattr(блок, "accept", None)
	return [часть.strip().lstrip(".").lower() for часть in (строка or "").split(",") if часть.strip()]


def расширение(имя: str) -> str:
	return имя.rsplit(".", 1)[-1].lower() if "." in имя else ""


def предел_байт() -> int:
	return int(настройка("artifact_file_max_mb", 10)) * 1024 * 1024


def проверить_ссылку(url: str) -> str:
	адрес = (url or "").strip()
	разбор = urlparse(адрес)
	if разбор.scheme not in ("http", "https") or not разбор.netloc:
		raise Отказ(НЕВЕРНАЯ_ССЫЛКА, "Ссылка должна начинаться с http:// или https://", url=url)
	return адрес


def из_base64(содержимое: str | None) -> bytes:
	"""Байты файла из base64; пустое или битое — пустые байты, и дальше
	`проверить_файл` откажет `artifact_file_missing`. Префикс `data:…;base64,`
	снимается: так файл отдаёт браузер."""
	import base64
	import binascii

	текст = (содержимое or "").strip()
	if текст.startswith("data:") and "," in текст:
		текст = текст.split(",", 1)[1]
	try:
		return base64.b64decode(текст, validate=True) if текст else b""
	except (binascii.Error, ValueError):
		return b""


def проверить_файл(блок, имя: str, данные: bytes, **где) -> str:
	"""Годится ли файл для блока; возвращает расширение."""
	if вид(блок) != ФАЙЛ:
		raise Отказ(ВИД_НЕ_ТОТ, "Этот блок не принимает файлы", kind=вид(блок), **где)
	if not данные:
		raise Отказ(ФАЙЛА_НЕТ, "Файл не передан или пуст", **где)
	тип = расширение(имя)
	можно = допустимые(блок)
	if можно and тип not in можно:
		raise Отказ(
			ФАЙЛ_НЕ_ТОГО_ТИПА,
			"Блок принимает файлы: " + ", ".join(можно),
			accept=можно,
			received=тип,
			**где,
		)
	if len(данные) > предел_байт():
		raise Отказ(
			ФАЙЛ_СЛИШКОМ_БОЛЬШОЙ,
			"Файл больше допустимого",
			max_bytes=предел_байт(),
			size=len(данные),
			**где,
		)
	return тип


def срез(данные: bytes, тип: str) -> str | None:
	"""Шапка и первые строки таблицы markdown-текстом; не таблица — `None`.

	Нечитаемый файл — тоже `None`, а не отказ: файл ученику нужен и без среза,
	а агент по его отсутствию скажет, что заглянуть внутрь не может.
	"""
	if тип not in ТАБЛИЦЫ:
		return None
	строк = int(настройка("artifact_preview_rows", 20))
	try:
		строки = _строки_csv(данные, строк) if тип == "csv" else _строки_xlsx(данные, строк)
	except Exception:
		frappe.log_error(title="Срез файла документа не построен")
		return None
	return _таблица(строки)


def _строки_csv(данные: bytes, предел: int) -> list[list[str]]:
	текст = данные.decode("utf-8-sig", errors="replace")
	try:
		диалект = csv.Sniffer().sniff(текст[:4096], delimiters=",;\t")
	except csv.Error:
		диалект = csv.excel
	строки = []
	for номер, строка in enumerate(csv.reader(io.StringIO(текст), диалект)):
		if номер > предел:
			break
		строки.append(строка)
	return строки


def _строки_xlsx(данные: bytes, предел: int) -> list[list[str]]:
	"""Первый лист: значение ячейки, а где значения нет — формула.

	`Why:` файл, сохранённый Excel или Google Таблицами, хранит посчитанные
	значения рядом с формулами — их агент и сверяет. У файла, собранного
	программой, значений нет, и формула лучше пустой клетки.
	"""
	from openpyxl import load_workbook

	значения = load_workbook(io.BytesIO(данные), read_only=True, data_only=True).worksheets[0]
	формулы = load_workbook(io.BytesIO(данные), read_only=True, data_only=False).worksheets[0]
	строки = []
	for номер, (ряд_значений, ряд_формул) in enumerate(
		zip(значения.iter_rows(values_only=True), формулы.iter_rows(values_only=True), strict=False)
	):
		if номер > предел:
			break
		строки.append(
			[
				"" if (з if з is not None else ф) is None else str(з if з is not None else ф)
				for з, ф in zip(ряд_значений, ряд_формул, strict=False)
			]
		)
	return строки


def _таблица(строки: list[list[str]]) -> str | None:
	строки = [с for с in строки if any((я or "").strip() for я in с)]
	if not строки:
		return None
	# Ширина — по последней непустой колонке. `Why:` Excel хранит
	# отформатированные, но пустые колонки, и на живом файле срез тянул за
	# шапкой два десятка безымянных пустых столбцов (#315).
	ширина = max(
		max((номер + 1 for номер, я in enumerate(с) if (я or "").strip()), default=0) for с in строки
	)
	строки = [с[:ширина] for с in строки]

	def ячейка(текст: str) -> str:
		текст = " ".join((текст or "").split()).replace("|", "\\|")
		return текст if len(текст) <= ШИРИНА_ЯЧЕЙКИ else текст[: ШИРИНА_ЯЧЕЙКИ - 1] + "…"

	def ряд(строка: list[str]) -> str:
		полная = list(строка) + [""] * (ширина - len(строка))
		return "| " + " | ".join(ячейка(я) for я in полная) + " |"

	шапка, *тело = строки
	return "\n".join([ряд(шапка), "|" + " --- |" * ширина, *(ряд(с) for с in тело)])


def сохранить_файл(документ_ученика, имя: str, данные: bytes):
	"""Приватный `File`, привязанный к документу ученика."""
	return frappe.get_doc(
		{
			"doctype": "File",
			"file_name": имя,
			"content": данные,
			"is_private": 1,
			"attached_to_doctype": "Agent Student Artifact",
			"attached_to_name": документ_ученика.name,
		}
	).insert(ignore_permissions=True)


def удалить_файл(имя_файла: str | None) -> None:
	if имя_файла and frappe.db.exists("File", имя_файла):
		frappe.delete_doc("File", имя_файла, ignore_permissions=True, force=True)


def сведения_о_файлах(имена: list[str]) -> dict[str, dict]:
	"""Что показать о файлах блоков — одним запросом."""
	if not имена:
		return {}
	return {
		запись.name: {
			"name": запись.file_name,
			"type": расширение(запись.file_name or ""),
			"size": запись.file_size or 0,
			"uploaded_at": запись.creation.isoformat() if запись.creation else None,
			"url": запись.file_url,
		}
		for запись in frappe.get_all(
			"File",
			filters={"name": ("in", имена)},
			fields=["name", "file_name", "file_size", "creation", "file_url"],
		)
	}
